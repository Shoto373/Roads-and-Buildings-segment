"""
Скрипт генерации отчета об изменениях в кодовой базе проекта (.docx).
Оформление строго в черном цвете, без цветных шрифтов, без ИИ-штампов и вводных фраз.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_margins(cell, top=100, bottom=100, left=160, right=160):
    """Установка внутренних отступов ячейки таблицы."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_code_box(doc, code_text):
    """Рамка для фрагментов кода с черным шрифтом."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)

    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F8F8"/>')
    cell._tc.get_or_add_tcPr().append(shading)

    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    set_cell_margins(cell, top=90, bottom=90, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    lines = code_text.strip().split('\n')
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if idx < len(lines) - 1:
            p.add_run('\n')

    spacing_p = doc.add_paragraph()
    spacing_p.paragraph_format.space_before = Pt(0)
    spacing_p.paragraph_format.space_after = Pt(3)


def add_note(doc, label, text):
    """Строгое текстовое примечание в черном цвете."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15

    run_label = p.add_run(label + ": ")
    run_label.bold = True
    run_label.font.color.rgb = RGBColor(0, 0, 0)

    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(0, 0, 0)


def add_divider(doc):
    """Горизонтальный разделитель."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="CCCCCC"/>'
        f'</w:pBdr>'
    )
    p._p.get_or_add_pPr().append(pBdr)


def add_footer_page_numbers(doc):
    """Центрированная нумерация страниц в нижнем колонтитуле."""
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        run._r.append(fldSimple)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document()

    # Поля страницы 1 дюйм
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Базовые стили
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].font.color.rgb = RGBColor(0, 0, 0)

    # Заголовок
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run('Отчет об изменениях в кодовой базе')
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run('Сегментация дорог и зданий на аэрофотоснимках (Roads & Buildings Segmentation)')
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(14)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(0, 0, 0)

    # =========================================================================
    # 1. Задачи доработки
    # =========================================================================
    h1 = doc.add_heading('1. Задачи доработки проекта', level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph('Перечень выполненных технических задач:')
    p.paragraph_format.space_after = Pt(2)

    tasks = [
        ("Оптимизация сегментации дорог: ", "устранение разрывов дорожной сети за счет замены функции потерь на Combo Loss."),
        ("Устранение дублирования кода: ", "объединение пяти независимых скриптов в модульную систему с общим CLI."),
        ("Смешанная точность (AMP): ", "поддержка FP16 через PyTorch AMP для снижения расхода видеопамяти на GPU."),
        ("Формат чекпоинтов: ", "переход с сериализации объектов модели (pickle) на словари весов (state_dict) с метаданными."),
        ("Контроль сходимости: ", "добавление Early Stopping и логирования метрик в TensorBoard."),
        ("Совместимость с Windows: ", "исправление ошибок кодировки cp1251 и отключение предупреждений OpenCV о TIFF-файлах.")
    ]
    for b_title, b_desc in tasks:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r1 = bp.add_run(b_title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0, 0, 0)
        r2 = bp.add_run(b_desc)
        r2.font.color.rgb = RGBColor(0, 0, 0)

    add_divider(doc)

    # =========================================================================
    # 2. Изменения в коде (Было / Стало)
    # =========================================================================
    h1 = doc.add_heading('2. Изменения в кодовой базе (Было / Стало)', level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    # 2.1
    h2 = doc.add_heading('2.1. Структура скриптов', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph('Было (разрозненные скрипты с дублированием кода):')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "train.py                 # Обучение только для зданий\n"
        "train_road.py            # Копия train.py для дорог\n"
        "trained_demo.py          # Инференс только для зданий\n"
        "trained_demo_road.py     # Инференс только для дорог\n"
        "generate_10_examples.py  # Скрипт генерации 10 примеров"
    )

    p = doc.add_paragraph('Стало (единая модульная структура):')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "config.py    # Единая конфигурация и CLI (--task road / building)\n"
        "train.py     # Единое обучение: AMP, Combo Loss, Early Stopping\n"
        "inference.py # Единый инференс: TTA, пакетная или одиночная обработка\n"
        "evaluate.py  # Расчет метрик (IoU, Dice, Precision, Recall, Acc) в JSON"
    )

    # 2.2
    h2 = doc.add_heading('2.2. Функция потерь для дорог', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph('Было:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(doc, "criterion = smp.losses.DiceLoss(mode='binary')")

    p = doc.add_paragraph('Стало:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "class ComboLoss(nn.Module):\n"
        "    def __init__(self, dice_weight=1.0, bce_weight=1.0):\n"
        "        super().__init__()\n"
        "        self.dice = smp.losses.DiceLoss(mode='binary', from_logits=True)\n"
        "        self.bce = nn.BCEWithLogitsLoss()\n"
        "        self.dice_weight = dice_weight\n"
        "        self.bce_weight = bce_weight\n\n"
        "    def forward(self, pred, target):\n"
        "        return self.dice_weight * self.dice(pred, target) + self.bce_weight * self.bce(pred, target)"
    )

    # 2.3
    h2 = doc.add_heading('2.3. Смешанная точность (AMP)', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph('Было:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "outputs = model(images)\n"
        "loss = criterion(outputs, masks)\n"
        "loss.backward()\n"
        "optimizer.step()"
    )

    p = doc.add_paragraph('Стало:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "scaler = torch.amp.GradScaler('cuda', enabled=use_amp)\n"
        "with torch.amp.autocast('cuda', enabled=use_amp):\n"
        "    outputs = model(images)\n"
        "    loss = criterion(outputs, masks)\n"
        "scaler.scale(loss).backward()\n"
        "scaler.step(optimizer)\n"
        "scaler.update()"
    )

    # 2.4
    h2 = doc.add_heading('2.4. Сохранение чекпоинтов', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph('Было:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(doc, "torch.save(model, './best_road_model.pth')")

    p = doc.add_paragraph('Стало:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(
        doc,
        "checkpoint = {\n"
        "    'epoch': epoch,\n"
        "    'state_dict': model.state_dict(),\n"
        "    'best_iou': best_iou_score,\n"
        "    'model_config': {'encoder': cfg.model.encoder, 'decoder': cfg.model.decoder},\n"
        "    'task': cfg.task\n"
        "}\n"
        "torch.save(checkpoint, save_path)"
    )

    add_note(doc, 'Примечание', 'Загрузчик весов в inference.py и evaluate.py поддерживает как новые state_dict, так и старые pickle-файлы.')

    # 2.5
    h2 = doc.add_heading('2.5. Мониторинг и ранняя остановка', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph('В train.py добавлены:')
    p.paragraph_format.space_after = Pt(2)
    p_b1 = doc.add_paragraph(style='List Bullet')
    p_b1.add_run('TensorBoard: ').bold = True
    p_b1.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p_b1.add_run('запись потерь и IoU (runs/road, runs/building).').font.color.rgb = RGBColor(0, 0, 0)
    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.add_run('Early Stopping: ').bold = True
    p_b2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p_b2.add_run('остановка обучения при отсутствии роста валидационного IoU (--patience).').font.color.rgb = RGBColor(0, 0, 0)
    p_b3 = doc.add_paragraph(style='List Bullet')
    p_b3.add_run('tqdm: ').bold = True
    p_b3.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p_b3.add_run('вывод текущего loss на каждом батче.').font.color.rgb = RGBColor(0, 0, 0)

    # 2.6
    h2 = doc.add_heading('2.6. Test-Time Augmentation (TTA)', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph('Усреднение предсказаний исходного кадра, горизонтального и вертикального отражений:')
    p.paragraph_format.space_after = Pt(2)
    add_code_box(doc, "python inference.py --task road --tta --num-examples 8")

    add_divider(doc)

    # =========================================================================
    # 3. Результаты тестирования
    # =========================================================================
    h1 = doc.add_heading('3. Результаты тестирования', level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    h2 = doc.add_heading('3.1. Сводная таблица метрик на тестовом наборе', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    metrics_data = [
        ["Метрика", "Сегментация дорог (Road, 49 снимков)", "Сегментация зданий (Building, 10 снимков)"],
        ["IoU (Intersection over Union)", "0.4872 ± 0.0603", "0.6039 ± 0.0351"],
        ["Dice / F1-Score", "0.6528 ± 0.0582", "0.7524 ± 0.0273"],
        ["Precision", "0.5704 ± 0.0492", "0.7908 ± 0.0268"],
        ["Recall (полнота)", "0.7730 ± 0.0991", "0.7187 ± 0.0385"],
        ["Accuracy (пиксельная)", "0.9626 ± 0.0209", "0.9143 ± 0.0340"],
    ]

    m_table = doc.add_table(rows=len(metrics_data), cols=3)
    m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_table.autofit = False

    col_widths = [Inches(2.5), Inches(2.0), Inches(2.0)]
    for row_idx, row in enumerate(m_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.text = metrics_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)

            p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
            if row_idx == 0:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F0F0"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(10)
            else:
                p.runs[0].font.size = Pt(10)
                if col_idx == 0:
                    p.runs[0].bold = True

            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                f'</w:tcBorders>'
            )
            cell._tc.get_or_add_tcPr().append(borders)
            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)

    spacing_p = doc.add_paragraph()
    spacing_p.paragraph_format.space_before = Pt(0)
    spacing_p.paragraph_format.space_after = Pt(6)

    h2 = doc.add_heading('3.2. Примеры сегментации дорог и зданий', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    # Дороги 1
    if os.path.exists('outputs/road_example_1.png'):
        doc.add_picture('outputs/road_example_1.png', width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run('Рис. 1: Сегментация дорог (Пример 1: Городская дорожная сеть)')
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(0, 0, 0)

    # Дороги 2
    if os.path.exists('outputs/road_example_2.png'):
        doc.add_picture('outputs/road_example_2.png', width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run('Рис. 2: Сегментация дорог (Пример 2: Магистрали и развязки)')
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(0, 0, 0)

    # Здания
    b_img = 'outputs/building_example_1.png' if os.path.exists('outputs/building_example_1.png') else 'outputs/example_1.png'
    if os.path.exists(b_img):
        doc.add_picture(b_img, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run('Рис. 3: Сегментация зданий (Демонстрационный пример малоэтажной застройки)')
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(0, 0, 0)

    add_divider(doc)

    # =========================================================================
    # 4. Команды запуска
    # =========================================================================
    h1 = doc.add_heading('4. Команды запуска', level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    h2 = doc.add_heading('4.1. Обучение', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    add_code_box(
        doc,
        "python train.py --task road --epochs 40 --batch-size 4 --loss combo\n"
        "python train.py --task building --epochs 40 --batch-size 4 --loss combo"
    )

    h2 = doc.add_heading('4.2. Инференс', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    add_code_box(
        doc,
        "python inference.py --task road --num-examples 8\n"
        "python inference.py --task road --tta --num-examples 8"
    )

    h2 = doc.add_heading('4.3. Оценка метрик', level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(3)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    add_code_box(
        doc,
        "python evaluate.py --task road\n"
        "python evaluate.py --task building"
    )

    # Нумерация страниц
    add_footer_page_numbers(doc)

    # Сохранение
    output_path = 'reports/Отчет_Изменения.docx'
    os.makedirs('reports', exist_ok=True)
    doc.save(output_path)
    print(f"[+] Successfully generated: {output_path}")


if __name__ == '__main__':
    main()
