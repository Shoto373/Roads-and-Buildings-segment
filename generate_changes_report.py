import sys
import subprocess
import os

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r"e:\WorkFOTON\Building-and-Road-Segmentation-from-Aerial-Images\reports\Отчет_Изменения.docx"
doc = docx.Document()

# Title
title = doc.add_heading("Отчет об изменениях: Обучение моделей сегментации", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("В ходе последних сессий была проделана огромная работа по обучению двух моделей для бинарной сегментации на основе архитектуры EfficientNet-B7 + UNet. Ниже представлен подробный пошаговый отчет о выполненных задачах для обоих датасетов.")

# Part 1: Buildings
doc.add_heading("Часть 1: Сегментация зданий (Massachusetts Buildings Dataset)", level=1)
doc.add_paragraph("В первой части мы настраивали пайплайн и обучали модель на датасете со зданиями (около 1.5 ГБ).")

doc.add_heading("Выполненные шаги:", level=2)
p1 = doc.add_paragraph(style='List Number')
p1.add_run("Подготовка датасета: ").bold = True
p1.add_run("Данные загружены и распакованы в папку notebooks/input/massachusetts-buildings-dataset.")

p2 = doc.add_paragraph(style='List Number')
p2.add_run("Настройка скрипта обучения: ").bold = True
p2.add_run("В файле train.py количество эпох увеличено до 30, а также установлен параметр num_workers=0 для стабильной работы DataLoader в среде Windows.")

p3 = doc.add_paragraph(style='List Number')
p3.add_run("Обучение: ").bold = True
p3.add_run("Модель обучалась исключительно на центральном процессоре (CPU). Благодаря оптимизациям и относительно небольшому размеру данных, процесс занял около 1 часа 15 минут.")

p4 = doc.add_paragraph(style='List Number')
p4.add_run("Оценка и визуализация: ").bold = True
p4.add_run("Достигнут высокий показатель Validation IoU (0.8505). Сгенерированы графики iou_score_plot.png и dice_loss_plot.png. Создана визуальная демонстрация предсказаний на тестовом снимке с помощью скрипта trained_demo.py.")

doc.add_paragraph("Примечание: Модель зданий показала отличную сходимость уже на 20-й эпохе, но продолжение до 30 эпох позволило зафиксировать стабильно низкий уровень функции потерь.", style='Intense Quote')


# Part 2: Roads
doc.add_heading("Часть 2: Сегментация дорог (Massachusetts Roads Dataset)", level=1)
doc.add_paragraph("Во второй части мы масштабировали наше решение на гораздо более объемный и сложный датасет (около 6 ГБ).")

doc.add_heading("Выполненные шаги:", level=2)
r1 = doc.add_paragraph(style='List Number')
r1.add_run("Скачивание данных: ").bold = True
r1.add_run("Через Kaggle CLI скачан массивный датасет дорог. Возникшая проблема со структурой папок (архив распаковался напрямую в notebooks/input) была оперативно исправлена с помощью скриптов PowerShell.")

r2 = doc.add_paragraph(style='List Number')
r2.add_run("Адаптация кода: ").bold = True
r2.add_run("Создан отдельный скрипт train_road.py. В нем классы изменены на ['background', 'road'], а пути сохранения весов и графиков обновлены (best_road_model.pth, road_iou_score_plot.png и т.д.).")

r3 = doc.add_paragraph(style='List Number')
r3.add_run("Ресурсоемкое обучение: ").bold = True
r3.add_run("Запущен процесс обучения на 30 эпох. Из-за гигантского объема данных (6 ГБ) и использования CPU обучение стало настоящим испытанием и продлилось почти 9 часов (с 16:50 до 01:46 следующего дня). Система успешно выдержала нагрузку без переполнения памяти, методично улучшая показатели качества.")

r4 = doc.add_paragraph(style='List Number')
r4.add_run("Оценка и визуализация: ").bold = True
r4.add_run("Достигнут высокий показатель Validation IoU (0.8968). Создан и запущен скрипт trained_demo_road.py для генерации графической демонстрации предсказаний дорожной сети.")

doc.add_paragraph("Важно: Обучение на процессоре таких тяжелых моделей занимает много времени. Использование графического ускорителя (GPU) позволило бы сократить этот процесс в разы.", style='Intense Quote')

# Conclusion
doc.add_heading("Обновление документации", level=1)
doc.add_paragraph("В рамках завершения задачи были обновлены ключевые документы проекта:")
doc.add_paragraph("1. README.md дополнен новыми секциями с результатами 30 эпох для обоих датасетов, графиками и описанием железа (CPU) и времени.", style='List Bullet')
doc.add_paragraph("2. Основной Отчет.docx перегенерирован с подробным описанием и всеми графиками.", style='List Bullet')
doc.add_paragraph("3. Все задачи в внутреннем трекере (task.md) успешно выполнены и закрыты.", style='List Bullet')

doc.save(doc_path)
print(f"Отчет об изменениях успешно сгенерирован: {doc_path}")
